"""file_processor — do something with a file the user dropped in (or named).

Two kinds of request:

* Understanding — summarize, describe, transcribe, explain, review, OCR,
  or any free-form instruction. The file itself goes to Gemini, which reads
  PDFs, images, audio, video and text natively; Office formats are turned
  into text locally first. Long answers are also saved next to the file.
* Transforming — resize/convert images, PDF → Word, filter/sort/convert
  tables, format JSON, trim/convert media (ffmpeg), list/extract archives.
  These are done locally and write a new file beside the original.
"""

from __future__ import annotations

import json
import mimetypes
import shutil
import subprocess
import tarfile
import tempfile
import time
import zipfile
from pathlib import Path

from toolkit import llm, runner
from toolkit.base import ToolContext, register, S, I, N, B

_INLINE_LIMIT = 18 * 1024 * 1024       # bigger files go through the Files API
_SAVE_ABOVE = 700                      # characters; longer answers are saved too
_NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)

_KIND = {}
for _kind, _exts in {
    "image": "jpg jpeg png gif webp bmp tif tiff heic",
    "pdf": "pdf",
    "word": "docx",
    "slides": "pptx",
    "table": "csv tsv xlsx xls ods",
    "json": "json",
    "audio": "mp3 wav m4a aac flac ogg opus wma",
    "video": "mp4 mov mkv avi webm wmv m4v",
    "archive": "zip tar gz tgz bz2 xz",
    "code": "py js ts jsx tsx java c cpp cs go rs rb php swift kt sh ps1 sql html css",
}.items():
    for _e in _exts.split():
        _KIND["." + _e] = _kind

# What the model is asked for each "understanding" action.
_ASKS = {
    "summarize": "Summarize this clearly: the main points first, then anything notable.",
    "describe": "Describe this in detail.",
    "analyze": "Analyse this thoroughly and point out what matters most.",
    "ocr": "Extract all of the text exactly as written, keeping the layout where it helps.",
    "extract_text": "Extract all of the text exactly as written.",
    "transcribe": "Transcribe all speech word for word. Mark speaker changes if they're clear.",
    "explain": "Walk through what this code does for someone who didn't write it.",
    "review": "Review this code: list real bugs and risks first, then worthwhile improvements.",
    "fix": "Return a corrected version of this file that fixes its bugs, then list what you changed.",
    "document": "Return this code with clear docstrings and comments added.",
    "test": "Write a set of unit tests covering this code.",
    "fix_text": "Correct the spelling, grammar and style of this text and return the corrected text.",
    "reformat": "Reformat this into a clean, well-structured document.",
    "to_bullet": "Turn this into a concise bullet-point summary.",
    "translate_hint": "Say what language this is in and summarize what it says.",
}


# --- helpers -----------------------------------------------------------------

def _beside(src: Path, label: str, ext: str | None = None) -> Path:
    out = src.with_name(f"{src.stem} ({label}){ext or src.suffix}")
    n = 2
    while out.exists():
        out = src.with_name(f"{src.stem} ({label} {n}){ext or src.suffix}")
        n += 1
    return out


def _office_text(path: Path) -> str:
    kind = _KIND.get(path.suffix.lower())
    if kind == "word":
        from docx import Document
        return "\n".join(p.text for p in Document(path).paragraphs)
    if kind == "slides":
        from pptx import Presentation
        slides = []
        for n, slide in enumerate(Presentation(path).slides, 1):
            texts = [s.text for s in slide.shapes if getattr(s, "has_text_frame", False) and s.text.strip()]
            slides.append(f"[Slide {n}]\n" + "\n".join(texts))
        return "\n\n".join(slides)
    if kind == "table":
        frame = _read_table(path)
        return f"{len(frame)} rows × {len(frame.columns)} columns.\n" + frame.head(200).to_csv(index=False)
    return path.read_text(encoding="utf-8", errors="replace")


def _file_part(path: Path):
    """The file as something Gemini can read: inline bytes, an uploaded file,
    or (for formats it can't open) extracted text."""
    kind = _KIND.get(path.suffix.lower())
    if kind in ("word", "slides", "table"):
        return _office_text(path)[:200_000]
    mime = mimetypes.guess_type(path.name)[0] or "text/plain"
    if kind in (None, "code", "json") or mime.startswith("text/"):
        return path.read_text(encoding="utf-8", errors="replace")[:200_000]
    if path.stat().st_size <= _INLINE_LIMIT:
        return llm.image_part(path.read_bytes(), mime)
    uploaded = llm.client().files.upload(file=str(path))
    while getattr(uploaded.state, "name", "") == "PROCESSING":
        time.sleep(2)
        uploaded = llm.client().files.get(name=uploaded.name)
    return uploaded


def understand(path: Path, request: str, save: bool) -> str:
    answer = llm.ask([_file_part(path), f"File name: {path.name}\n\n{request}"])
    if save and len(answer) > _SAVE_ABOVE:
        out = _beside(path, "notes", ".txt")
        out.write_text(answer, encoding="utf-8")
        return f"{answer[:600]}…\n\n(Full answer saved to {out})"
    return answer


def _ffmpeg(*args: str, timeout: int = 1800) -> None:
    if not shutil.which("ffmpeg"):
        raise RuntimeError("ffmpeg isn't installed — it's needed for audio/video editing")
    done = subprocess.run(["ffmpeg", "-y", "-hide_banner", "-loglevel", "error", *args],
                          capture_output=True, text=True, timeout=timeout, creationflags=_NO_WINDOW)
    if done.returncode != 0:
        raise RuntimeError(done.stderr.strip()[-300:] or "ffmpeg failed")


def _read_table(path: Path):
    import pandas as pd
    if path.suffix.lower() in (".csv", ".tsv"):
        return pd.read_csv(path, sep="\t" if path.suffix.lower() == ".tsv" else ",", encoding_errors="replace")
    return pd.read_excel(path)


def _write_table(frame, src: Path, label: str, fmt: str) -> Path:
    fmt = fmt.lower().lstrip(".") or "csv"
    out = _beside(src, label, "." + ("xlsx" if fmt in ("excel", "xlsx") else fmt))
    if out.suffix == ".xlsx":
        frame.to_excel(out, index=False)
    elif out.suffix == ".json":
        frame.to_json(out, orient="records", indent=2, force_ascii=False)
    else:
        frame.to_csv(out, index=False)
    return out


# --- local transforms --------------------------------------------------------

def _image(path: Path, action: str, a: dict) -> str | None:
    from PIL import Image
    if action not in ("resize", "convert", "compress", "info"):
        return None
    img = Image.open(path)
    if action == "info":
        return f"{img.format} image, {img.width}×{img.height}, {img.mode}, {path.stat().st_size / 1024:.0f} KB."
    if action == "resize":
        w, h = int(a.get("width") or 0), int(a.get("height") or 0)
        scale = float(a.get("scale") or 0)
        if scale:
            w, h = round(img.width * scale), round(img.height * scale)
        elif w and not h:
            h = round(img.height * w / img.width)
        elif h and not w:
            w = round(img.width * h / img.height)
        if not (w and h):
            return "Give a width, a height, or a scale factor."
        out = _beside(path, f"{w}x{h}")
        img.resize((w, h), Image.LANCZOS).save(out)
        return f"Resized to {w}×{h}: {out}"
    if action == "convert":
        fmt = (a.get("format") or "png").lower().lstrip(".")
        fmt = "jpeg" if fmt == "jpg" else fmt
        out = _beside(path, "converted", "." + ("jpg" if fmt == "jpeg" else fmt))
        (img.convert("RGB") if fmt in ("jpeg", "bmp") else img).save(out, fmt.upper())
        return f"Saved as {out}"
    quality = max(10, min(int(a.get("quality") or 75), 95))
    out = _beside(path, "smaller", ".jpg")
    img.convert("RGB").save(out, "JPEG", quality=quality, optimize=True)
    return f"Compressed {path.stat().st_size // 1024} KB → {out.stat().st_size // 1024} KB: {out}"


def _pdf(path: Path, action: str, a: dict) -> str | None:
    if action == "info":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return f"PDF with {len(pdf.pages)} pages, {path.stat().st_size / 1024:.0f} KB."
    if action in ("extract_text", "to_word"):
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            pages = [(p.extract_text() or "") for p in pdf.pages]
        text = "\n\n".join(pages).strip()
        if not text:
            return None   # scanned PDF — let the model read it instead
        if action == "extract_text":
            out = _beside(path, "text", ".txt")
            out.write_text(text, encoding="utf-8")
            return f"Extracted {len(text)} characters to {out}"
        from docx import Document
        doc = Document()
        for block in text.split("\n\n"):
            doc.add_paragraph(block.strip())
        out = _beside(path, "Word", ".docx")
        doc.save(out)
        return f"Converted to Word: {out}"
    return None


def _table(path: Path, action: str, a: dict) -> str | None:
    if action not in ("info", "stats", "filter", "sort", "convert", "to_csv", "to_excel", "to_json"):
        return None
    frame = _read_table(path)
    if action == "info":
        return f"{len(frame)} rows × {len(frame.columns)} columns: {', '.join(map(str, frame.columns))}"
    if action == "stats":
        return frame.describe(include="all").transpose().to_string()[:3500]
    if action.startswith("to_") or action == "convert":
        fmt = action[3:] if action.startswith("to_") else (a.get("format") or "csv")
        return f"Saved as {_write_table(frame, path, 'converted', fmt)}"
    column = a.get("column") or ""
    if column not in frame.columns:
        return f"There's no column {column!r}. The columns are {', '.join(map(str, frame.columns))}."
    if action == "sort":
        out = _write_table(frame.sort_values(column, ascending=a.get("ascending", True) is not False),
                           path, f"sorted by {column}", path.suffix)
        return f"Sorted by {column}: {out}"
    value, cond = a.get("value"), (a.get("condition") or "equals").lower()
    col = frame[column]
    if cond == "contains":
        mask = col.astype(str).str.contains(str(value), case=False, na=False)
    elif cond in ("gt", "lt"):
        import pandas as pd
        nums = pd.to_numeric(col, errors="coerce")
        mask = nums > float(value) if cond == "gt" else nums < float(value)
    else:
        mask = col.astype(str) == str(value)
    out = _write_table(frame[mask], path, "filtered", path.suffix)
    return f"{int(mask.sum())} of {len(frame)} rows match: {out}"


def _json(path: Path, action: str, a: dict) -> str | None:
    if action not in ("validate", "format", "to_csv", "info"):
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as err:
        return f"Not valid JSON: {err.msg} at line {err.lineno}, column {err.colno}."
    if action in ("validate", "info"):
        size = len(data) if isinstance(data, (list, dict)) else 1
        return f"Valid JSON — top level is a {type(data).__name__} with {size} entr{'y' if size == 1 else 'ies'}."
    if action == "format":
        out = _beside(path, "formatted")
        out.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
        return f"Pretty-printed to {out}"
    import pandas as pd
    out = _beside(path, "table", ".csv")
    pd.json_normalize(data if isinstance(data, list) else [data]).to_csv(out, index=False)
    return f"Flattened to CSV: {out}"


def _media(path: Path, action: str, a: dict) -> str | None:
    kind = _KIND.get(path.suffix.lower())
    if action == "info":
        if not shutil.which("ffprobe"):
            return f"{kind} file, {path.stat().st_size / 1_048_576:.1f} MB (install ffmpeg for details)."
        probe = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration:stream=codec_type,width,height",
                                "-of", "json", str(path)], capture_output=True, text=True, creationflags=_NO_WINDOW)
        info = json.loads(probe.stdout or "{}")
        secs = float(info.get("format", {}).get("duration") or 0)
        video = next((s for s in info.get("streams", []) if s.get("codec_type") == "video"), None)
        size = f", {video['width']}×{video['height']}" if video and video.get("width") else ""
        return f"{kind}, {int(secs // 60)}m {int(secs % 60)}s{size}, {path.stat().st_size / 1_048_576:.1f} MB."
    if action == "trim":
        start, end = str(a.get("start") or "0"), a.get("end")
        out = _beside(path, "trimmed")
        _ffmpeg("-ss", start, *(["-to", str(end)] if end else []), "-i", str(path), "-c", "copy", str(out))
        return f"Trimmed copy saved to {out}"
    if action == "convert":
        fmt = (a.get("format") or ("mp3" if kind == "audio" else "mp4")).lstrip(".")
        out = _beside(path, "converted", "." + fmt)
        _ffmpeg("-i", str(path), str(out))
        return f"Converted: {out}"
    if action == "extract_audio" and kind == "video":
        out = _beside(path, "audio", ".mp3")
        _ffmpeg("-i", str(path), "-vn", "-q:a", "2", str(out))
        return f"Audio saved to {out}"
    if action == "extract_frame" and kind == "video":
        stamp = str(a.get("timestamp") or "00:00:01")
        out = _beside(path, f"frame {stamp.replace(':', '.')}", ".jpg")
        _ffmpeg("-ss", stamp, "-i", str(path), "-frames:v", "1", str(out))
        return f"Frame saved to {out}"
    if action == "compress" and kind == "video":
        crf = max(18, min(int(a.get("quality") or 28), 40))
        out = _beside(path, "smaller", ".mp4")
        _ffmpeg("-i", str(path), "-c:v", "libx264", "-crf", str(crf), "-preset", "medium", "-c:a", "aac", str(out))
        return f"Compressed {path.stat().st_size >> 20} MB → {out.stat().st_size >> 20} MB: {out}"
    return None


def _archive(path: Path, action: str, a: dict) -> str | None:
    if action not in ("list", "extract", "info"):
        return None
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            if action != "extract":
                return f"{len(names)} entries:\n" + "\n".join(names[:60])
            dest = Path(a.get("destination") or path.with_suffix(""))
            for name in names:   # refuse entries that would land outside dest
                if not (dest / name).resolve().is_relative_to(dest.resolve()):
                    return f"Refusing to extract — the archive contains an unsafe path ({name})."
            z.extractall(dest)
            return f"Extracted {len(names)} entries to {dest}"
    if tarfile.is_tarfile(path):
        with tarfile.open(path) as t:
            if action != "extract":
                names = t.getnames()
                return f"{len(names)} entries:\n" + "\n".join(names[:60])
            dest = Path(a.get("destination") or path.with_suffix(""))
            t.extractall(dest, filter="data")
            return f"Extracted to {dest}"
    return "I can only open .zip and .tar archives."


_TRANSFORMS = {"image": _image, "pdf": _pdf, "table": _table, "json": _json,
               "audio": _media, "video": _media, "archive": _archive}


def _default_action(kind: str | None) -> str:
    return {"image": "describe", "audio": "transcribe", "video": "summarize",
            "archive": "list", "code": "explain", "table": "analyze"}.get(kind or "", "summarize")


@register(
    "file_processor",
    "Does something with a file the user uploaded or dropped on the window (or a "
    "named file): summarize/describe/analyze/transcribe/OCR/explain/review it, or "
    "transform it — resize/convert/compress images, PDF to Word or text, "
    "filter/sort/convert CSV and Excel, validate/format JSON, trim/convert media, "
    "list/extract archives. It is the way to act on an attached file; with "
    "file_path left empty it works on the file attached most recently.",
    {
        "file_path": S("Path to the file; leave empty for the file currently uploaded"),
        "action": S("summarize | describe | analyze | ocr | extract_text | transcribe | explain | review | "
                    "fix | document | test | fix_text | reformat | to_bullet | translate_hint | info | resize | "
                    "convert | compress | to_word | stats | filter | sort | to_csv | to_excel | to_json | "
                    "validate | format | trim | extract_audio | extract_frame | list | extract | run"),
        "instruction": S("A free-form request about the file, e.g. 'find every email address'"),
        "format": S("Target format for convert, e.g. png, mp3, xlsx"),
        "width": I("Resize width"), "height": I("Resize height"), "scale": N("Resize factor, e.g. 0.5"),
        "quality": I("Compression quality"),
        "start": S("Where the kept part begins, as seconds or HH:MM:SS"), "end": S("Where it ends, same format"),
        "timestamp": S("Video frame time HH:MM:SS"),
        "column": S("Table column for filter/sort"), "value": S("Filter value"),
        "condition": S("How to compare the column with value: equals, contains, gt (greater than) or lt (less than)"),
        "ascending": B("Sort ascending (default true)"),
        "save": B("Also save long answers to a text file (default true)"),
        "destination": S("Folder to extract an archive into"),
    },
)
def file_processor(args: dict, ctx: ToolContext) -> str:
    raw = (args.get("file_path") or "").strip() or (ctx.attached_file or "")
    if not raw:
        return "Which file? Drop one on the window or give me its path."
    path = Path(raw).expanduser()
    if not path.is_file():
        return f"There's no file at {path}."
    kind = _KIND.get(path.suffix.lower())
    action = (args.get("action") or "").strip().lower()
    instruction = (args.get("instruction") or "").strip()
    ctx.log(f"[documents] {path.name} → {action or instruction[:30] or 'auto'}")

    try:
        if action == "run" and kind == "code":
            ask = getattr(ctx.ui, "confirm_action", None)
            if not (callable(ask) and ask(f"Run {path.name}? This executes real code on your computer.")):
                return "Not run — the user didn't confirm."
            return runner.run_file(path).summary()
        transform = _TRANSFORMS.get(kind or "")
        if action and transform:
            done = transform(path, action, args)
            if done is not None:
                return done
        if action == "fix" and kind not in ("code", "json"):
            action = "fix_text"
        request = instruction or _ASKS.get(action or _default_action(kind)) or f"{action.replace('_', ' ')} this file."
        return understand(path, request, save=args.get("save", True) is not False)
    except Exception as err:
        return f"Working on {path.name} failed: {err}"
